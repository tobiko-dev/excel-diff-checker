"""Core Excel diff engine.

Compares two ``.xlsx`` workbooks and reports which sheets, rows, columns and
cells differ. Two row-alignment strategies are supported:

* **positional** (default): cell (r, c) in file1 is compared with cell (r, c)
  in file2. Simple and fast, but inserting a row near the top makes everything
  below it look "changed".
* **keyed**: rows are matched by the value of a key column (identified by its
  header name). This survives inserted/deleted/reordered rows and answers the
  question "which records changed?" far more usefully.
"""

from __future__ import annotations

from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .model import CellDiff, SheetDiff, WorkbookDiff


def compare_files(
    file1: str,
    file2: str,
    *,
    sheets: list[str] | None = None,
    key: str | None = None,
    header_row: int = 1,
    ignore_empty: bool = True,
) -> WorkbookDiff:
    """Compare two Excel workbooks.

    Args:
        file1: Path to the first (baseline) workbook.
        file2: Path to the second (comparison) workbook.
        sheets: Restrict comparison to these sheet names. ``None`` compares
            every sheet present in either file.
        key: Column header to use for keyed row matching. When provided, rows
            are aligned by their value in this column instead of by position.
        header_row: 1-indexed row that holds column headers. Only used to
            resolve ``key`` and to label columns in keyed mode.
        ignore_empty: Treat ``None`` and empty string as equal so a blank cell
            never counts as a difference against another blank cell.

    Returns:
        A :class:`WorkbookDiff` describing every difference found.
    """
    wb1 = load_workbook(file1, data_only=True, read_only=True)
    wb2 = load_workbook(file2, data_only=True, read_only=True)
    try:
        names1 = wb1.sheetnames
        names2 = wb2.sheetnames
        common = [n for n in names1 if n in names2]
        added_sheets = [n for n in names2 if n not in names1]
        removed_sheets = [n for n in names1 if n not in names2]

        if sheets is not None:
            wanted = set(sheets)
            common = [n for n in common if n in wanted]
            added_sheets = [n for n in added_sheets if n in wanted]
            removed_sheets = [n for n in removed_sheets if n in wanted]

        result = WorkbookDiff(
            file1=file1,
            file2=file2,
            added_sheets=added_sheets,
            removed_sheets=removed_sheets,
        )

        for name in common:
            grid1 = _read_grid(wb1[name])
            grid2 = _read_grid(wb2[name])
            if key:
                sheet_diff = _diff_sheet_keyed(
                    name, grid1, grid2, key, header_row, ignore_empty
                )
            else:
                sheet_diff = _diff_sheet_positional(
                    name, grid1, grid2, ignore_empty
                )
            result.sheet_diffs.append(sheet_diff)

        return result
    finally:
        wb1.close()
        wb2.close()


def _read_grid(ws) -> list[list[Any]]:
    """Read a worksheet into a list-of-rows grid of cell values."""
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    # openpyxl's read-only mode can emit trailing all-None rows; trim them so
    # dimensions reflect real content.
    while grid and _row_is_empty(grid[-1]):
        grid.pop()
    return grid


def _row_is_empty(row: list[Any]) -> bool:
    return all(v is None for v in row)


def _values_equal(a: Any, b: Any, ignore_empty: bool) -> bool:
    if ignore_empty:
        a_empty = a is None or a == ""
        b_empty = b is None or b == ""
        if a_empty and b_empty:
            return True
    return a == b


def _cell(grid: list[list[Any]], r: int, c: int) -> Any:
    """Return the value at 0-indexed (r, c), or None if out of range."""
    if r < len(grid) and c < len(grid[r]):
        return grid[r][c]
    return None


def _dims(grid: list[list[Any]]) -> tuple[int, int]:
    max_row = len(grid)
    max_col = max((len(r) for r in grid), default=0)
    return max_row, max_col


# --------------------------------------------------------------------------- #
# Positional comparison
# --------------------------------------------------------------------------- #
def _diff_sheet_positional(
    name: str,
    grid1: list[list[Any]],
    grid2: list[list[Any]],
    ignore_empty: bool,
) -> SheetDiff:
    rows1, cols1 = _dims(grid1)
    rows2, cols2 = _dims(grid2)
    diff = SheetDiff(
        name=name,
        status="identical",
        dims_old=(rows1, cols1),
        dims_new=(rows2, cols2),
    )

    max_rows = max(rows1, rows2)
    max_cols = max(cols1, cols2)

    for r in range(max_rows):
        for c in range(max_cols):
            v1 = _cell(grid1, r, c)
            v2 = _cell(grid2, r, c)
            if not _values_equal(v1, v2, ignore_empty):
                excel_row = r + 1
                excel_col = c + 1
                diff.cell_diffs.append(
                    CellDiff(
                        sheet=name,
                        row=excel_row,
                        column=excel_col,
                        column_letter=get_column_letter(excel_col),
                        old_value=v1,
                        new_value=v2,
                    )
                )
                diff.changed_rows.add(excel_row)
                diff.changed_columns.add(excel_col)

    if rows2 > rows1:
        diff.added_rows = list(range(rows1 + 1, rows2 + 1))
    elif rows1 > rows2:
        diff.removed_rows = list(range(rows2 + 1, rows1 + 1))

    if cols2 > cols1:
        diff.added_columns = list(range(cols1 + 1, cols2 + 1))
    elif cols1 > cols2:
        diff.removed_columns = list(range(cols2 + 1, cols1 + 1))

    if (
        diff.cell_diffs
        or diff.added_rows
        or diff.removed_rows
        or diff.added_columns
        or diff.removed_columns
    ):
        diff.status = "modified"
    return diff


# --------------------------------------------------------------------------- #
# Keyed comparison
# --------------------------------------------------------------------------- #
def _diff_sheet_keyed(
    name: str,
    grid1: list[list[Any]],
    grid2: list[list[Any]],
    key: str,
    header_row: int,
    ignore_empty: bool,
) -> SheetDiff:
    rows1, cols1 = _dims(grid1)
    rows2, cols2 = _dims(grid2)
    diff = SheetDiff(
        name=name,
        status="identical",
        dims_old=(rows1, cols1),
        dims_new=(rows2, cols2),
    )

    hdr_idx = header_row - 1
    headers1 = grid1[hdr_idx] if hdr_idx < len(grid1) else []
    headers2 = grid2[hdr_idx] if hdr_idx < len(grid2) else []

    key1 = _find_header(headers1, key)
    key2 = _find_header(headers2, key)
    if key1 is None or key2 is None:
        raise ValueError(
            f"Key column {key!r} not found in header row {header_row} "
            f"of sheet {name!r}."
        )

    # Map header name -> column index for each file, then compare the columns
    # that exist in both by name.
    cols_by_name1 = _headers_to_map(headers1)
    cols_by_name2 = _headers_to_map(headers2)
    common_cols = [h for h in cols_by_name1 if h in cols_by_name2 and h != key]
    added_cols = [h for h in cols_by_name2 if h not in cols_by_name1]
    removed_cols = [h for h in cols_by_name1 if h not in cols_by_name2]
    diff.added_columns = [cols_by_name2[h] + 1 for h in added_cols]
    diff.removed_columns = [cols_by_name1[h] + 1 for h in removed_cols]

    index1 = _index_by_key(grid1, key1, header_row)
    index2 = _index_by_key(grid2, key2, header_row)

    for key_val, r1 in index1.items():
        if key_val not in index2:
            diff.removed_rows.append(r1 + 1)

    for key_val, r2 in index2.items():
        if key_val not in index1:
            diff.added_rows.append(r2 + 1)

    for key_val, r1 in index1.items():
        r2 = index2.get(key_val)
        if r2 is None:
            continue
        for col_name in common_cols:
            c1 = cols_by_name1[col_name]
            c2 = cols_by_name2[col_name]
            v1 = _cell(grid1, r1, c1)
            v2 = _cell(grid2, r2, c2)
            if not _values_equal(v1, v2, ignore_empty):
                excel_row = r2 + 1
                excel_col = c2 + 1
                diff.cell_diffs.append(
                    CellDiff(
                        sheet=name,
                        row=excel_row,
                        column=excel_col,
                        column_letter=get_column_letter(excel_col),
                        old_value=v1,
                        new_value=v2,
                        row_key=key_val,
                        column_name=col_name,
                    )
                )
                diff.changed_rows.add(excel_row)
                diff.changed_columns.add(excel_col)

    diff.added_rows.sort()
    diff.removed_rows.sort()

    if (
        diff.cell_diffs
        or diff.added_rows
        or diff.removed_rows
        or diff.added_columns
        or diff.removed_columns
    ):
        diff.status = "modified"
    return diff


def _find_header(headers: list[Any], name: str) -> int | None:
    for i, h in enumerate(headers):
        if h == name:
            return i
    return None


def _headers_to_map(headers: list[Any]) -> dict[Any, int]:
    """Map header value -> column index, keeping the first occurrence."""
    mapping: dict[Any, int] = {}
    for i, h in enumerate(headers):
        if h is None or h == "":
            continue
        if h not in mapping:
            mapping[h] = i
    return mapping


def _index_by_key(
    grid: list[list[Any]], key_col: int, header_row: int
) -> dict[Any, int]:
    """Map key value -> row index for data rows below the header.

    On duplicate keys the first row wins, matching typical spreadsheet lookup
    behaviour.
    """
    index: dict[Any, int] = {}
    for r in range(header_row, len(grid)):
        val = _cell(grid, r, key_col)
        if val is None or val == "":
            continue
        if val not in index:
            index[val] = r
    return index
