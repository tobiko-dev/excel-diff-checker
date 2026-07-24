"""Tests for the Excel diff engine."""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from excel_diff import compare_files  # noqa: E402


def _write(path, rows, sheet_name="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(list(row))
    wb.save(path)


def test_identical_files(tmp_path):
    data = [["Name", "Qty"], ["Apple", 3], ["Pear", 5]]
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, data)
    _write(b, data)

    diff = compare_files(str(a), str(b))
    assert not diff.has_differences
    assert diff.sheet_diffs[0].status == "identical"


def test_changed_cell_positional(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, [["Name", "Qty"], ["Apple", 3], ["Pear", 5]])
    _write(b, [["Name", "Qty"], ["Apple", 3], ["Pear", 9]])

    diff = compare_files(str(a), str(b))
    sheet = diff.sheet_diffs[0]
    assert sheet.status == "modified"
    assert sheet.changed_rows == {3}
    assert sheet.changed_columns == {2}
    assert len(sheet.cell_diffs) == 1
    cell = sheet.cell_diffs[0]
    assert (cell.row, cell.column_letter) == (3, "B")
    assert cell.old_value == 5
    assert cell.new_value == 9


def test_added_and_removed_rows(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, [["Name", "Qty"], ["Apple", 3]])
    _write(b, [["Name", "Qty"], ["Apple", 3], ["Pear", 5]])

    diff = compare_files(str(a), str(b))
    sheet = diff.sheet_diffs[0]
    assert sheet.added_rows == [3]
    assert sheet.removed_rows == []


def test_added_column(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, [["Name"], ["Apple"]])
    _write(b, [["Name", "Qty"], ["Apple", 3]])

    diff = compare_files(str(a), str(b))
    sheet = diff.sheet_diffs[0]
    assert sheet.added_columns == [2]


def test_added_and_removed_sheets(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    wb = Workbook()
    wb.active.title = "Common"
    wb.active.append(["x"])
    wb.create_sheet("OnlyInA").append(["x"])
    wb.save(a)

    wb2 = Workbook()
    wb2.active.title = "Common"
    wb2.active.append(["x"])
    wb2.create_sheet("OnlyInB").append(["x"])
    wb2.save(b)

    diff = compare_files(str(a), str(b))
    assert diff.removed_sheets == ["OnlyInA"]
    assert diff.added_sheets == ["OnlyInB"]


def test_values_equal_empty_handling():
    # Excel/openpyxl normalizes an empty-string cell to None on read, so the
    # ignore_empty flag governs how None and "" compare at the value level.
    from excel_diff.diff import _values_equal

    assert _values_equal(None, "", ignore_empty=True)
    assert not _values_equal(None, "", ignore_empty=False)
    assert _values_equal(None, None, ignore_empty=False)
    assert not _values_equal(None, "x", ignore_empty=True)


def test_keyed_matching_survives_reorder(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    # Same records, different order, one value changed for id=2.
    _write(a, [["id", "name", "qty"], [1, "Apple", 3], [2, "Pear", 5]])
    _write(b, [["id", "name", "qty"], [2, "Pear", 9], [1, "Apple", 3]])

    # Positional would flag many cells; keyed should flag exactly one.
    diff = compare_files(str(a), str(b), key="id")
    sheet = diff.sheet_diffs[0]
    assert len(sheet.cell_diffs) == 1
    cell = sheet.cell_diffs[0]
    assert cell.row_key == 2
    assert cell.column_name == "qty"
    assert cell.old_value == 5
    assert cell.new_value == 9


def test_keyed_added_removed_rows(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, [["id", "name"], [1, "Apple"], [2, "Pear"]])
    _write(b, [["id", "name"], [1, "Apple"], [3, "Plum"]])

    diff = compare_files(str(a), str(b), key="id")
    sheet = diff.sheet_diffs[0]
    # id=2 removed, id=3 added
    assert sheet.removed_rows == [3]  # row 3 in file1 held id=2
    assert sheet.added_rows == [3]  # row 3 in file2 holds id=3


def test_keyed_missing_key_raises(tmp_path):
    a = tmp_path / "a.xlsx"
    b = tmp_path / "b.xlsx"
    _write(a, [["id", "name"], [1, "Apple"]])
    _write(b, [["id", "name"], [1, "Apple"]])

    import pytest

    with pytest.raises(ValueError):
        compare_files(str(a), str(b), key="nonexistent")
